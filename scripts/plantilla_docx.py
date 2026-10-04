#!/usr/bin/env python3
# Copyright 2026 lianfixx and contributors. SPDX-License-Identifier: Apache-2.0
"""Genera una plantilla institucional EN BLANCO. No emite un diagnostico de cliente."""
from __future__ import annotations
import argparse
from pathlib import Path


def generate(output: Path) -> None:
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc:
        raise RuntimeError('Instala la dependencia opcional: python3 -m pip install -r requirements-documentos.txt') from exc
    doc = Document()
    doc.core_properties.author = "lianfixx y colaboradores; plantilla MILLA Asesorías"
    doc.core_properties.subject = "Plantilla en blanco Apache-2.0; sin afiliación ni dictamen aprobado"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.8), Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.header_distance = sec.footer_distance = Inches(.3)
    blue, pale, gray = '17365D', 'EAF1F8', '667085'
    normal = doc.styles['Normal']; normal.font.name = 'Arial'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    for name, size in [('Title',22),('Heading 1',15),('Heading 2',11)]:
        style = doc.styles[name]; style.font.name = 'Arial'; style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(blue)
        style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run('[CLIENTE] | MA-[INICIALES]/[NN]'); run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string(gray)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run('CONFIDENCIAL · PLANTILLA SIN EMITIR | Página '); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(gray)
    def field(paragraph, code):
        run = paragraph.add_run(); item = OxmlElement('w:fldSimple'); item.set(qn('w:instr'), code); run._r.addnext(item)
    field(footer, 'PAGE'); footer.add_run(' de '); field(footer, 'NUMPAGES')
    def shade(cell, fill):
        pr = cell._tc.get_or_add_tcPr(); el = OxmlElement('w:shd'); el.set(qn('w:fill'), fill); pr.append(el)
    def text(value, small=False):
        p = doc.add_paragraph(value)
        if small:
            for r in p.runs: r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(gray)
        return p
    def table(headers, rows, widths=None):
        t = doc.add_table(rows=1, cols=len(headers)); t.autofit = False
        for i, label in enumerate(headers):
            cell = t.rows[0].cells[i]; cell.text = label; shade(cell, blue)
            for r in cell.paragraphs[0].runs: r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(9)
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for i, value in enumerate(row):
                cells[i].text=value
                if i==0: shade(cells[i], pale)
                for p in cells[i].paragraphs:
                    p.paragraph_format.space_after=Pt(5); p.paragraph_format.space_before=Pt(5)
                    for r in p.runs: r.font.size=Pt(9)
        for row in t.rows:
            pr=row._tr.get_or_add_trPr(); no=OxmlElement('w:cantSplit'); pr.append(no)
            if widths:
                for i, cell in enumerate(row.cells): cell.width=Inches(widths[i])
        text('')
        return t
    def box(title, body):
        t=doc.add_table(rows=1,cols=1); cell=t.cell(0,0); shade(cell,pale)
        p=cell.paragraphs[0]; r=p.add_run(title+' '); r.bold=True; r.font.color.rgb=RGBColor.from_string(blue)
        p.add_run(body)
        for p in cell.paragraphs:
            p.paragraph_format.space_before=Pt(8); p.paragraph_format.space_after=Pt(8)
        text('')
    def page(title):
        doc.add_page_break(); doc.add_heading(title,1)
    text('MILLA ABOGADOS · DOCUMENTO BASE PARA REVISIÓN', True)
    p=doc.add_paragraph('DIAGNÓSTICO JURÍDICO\nPRELIMINAR', 'Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p=text('[Materia y objetivo específico del asunto]',True); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    table(['Dato','Información'],[['Cliente','[Nombre verificado]'],['Materia y posición','[Materia / posición jurídica]'],['Ámbito probable','[País, entidad, autoridad y pendientes]'],['Etapa actual','[Situación real]'],['Corte fáctico','[Fecha de información del asunto]'],['Consulta normativa','[Fecha real de verificación]'],['Control','[Versión / borrador / folio al emitir]']], [1.65,5.35])
    box('Naturaleza y alcance.', 'Informe preliminar sujeto a los hechos y documentos disponibles. No constituye una resolución ni garantía de resultado. Debe identificar las cuestiones pendientes y contar con revisión profesional antes de su entrega.')
    text('PLANTILLA EN BLANCO. Los campos entre corchetes deben adaptarse. Esta edición no contiene datos, honorarios ni conclusiones de un cliente.',True)
    page('Contenido y uso del documento')
    for title in ['1. Objeto y síntesis ejecutiva','2. Hechos y valoración jurídica','3. Ruta de MILLA ABOGADOS','4. Prueba y documentación necesaria','5. Tiempos y condiciones','6. Propuesta de honorarios','7. Próximos pasos y revisión']:
        text(title)
    box('Lectura del diagnóstico.', 'La situación se explica primero; después se presentan alternativas, riesgos y acciones. Los detalles técnicos y documentos reservados pueden conservarse en una carpeta interna separada, sin ocultar al cliente incertidumbres relevantes.')
    text('Adaptar el índice a la materia. No forzar conciliación, juicio, peritos o recursos que no correspondan al caso. Actualizar la numeración al generar la versión final.',True)
    page('1. Objeto y síntesis ejecutiva')
    text('[Explicar qué se revisó, qué se pretende resolver, recomendación preliminar y principal condición o límite.]')
    doc.add_heading('2. Hechos y valoración jurídica',1)
    table(['Hecho / problema','Fuente y estado','Valoración / pendiente'],[['[Hecho relevante]','[Documento / manifestación / fecha]','[Qué permite concluir y qué no]'],['[Contradicción o falta]','[Fuente]','[Cómo modifica la decisión]']], [2,2.2,2.8])
    text('[Por cada problema: hecho y evidencia → norma verificada → aplicación al caso → objeción o riesgo → conclusión provisional.]')
    box('Punto que debe confirmarse.', '[Identificar la información cuya ausencia impide una conclusión o actuación segura y la forma de obtenerla.]')
    page('3. Ruta de MILLA ABOGADOS')
    table(['Fase','Acciones de la firma','Participación del cliente','Resultado o decisión'],[['[Fase pertinente]','[Trabajo concreto]','[Documentos / autorización]','[Qué permite avanzar]'],['[Alternativa]','[Acción]','[Decisión]','[Condición de cambio]']], [1.1,2.1,1.8,2.0])
    doc.add_heading('4. Prueba y documentación necesaria',1)
    table(['Documento / prueba','Utilidad y estado','Prioridad / obtención'],[['[Elemento necesario]','[Hecho que ayuda a acreditar]','[Prioridad y responsable]'],['[Especialista, sólo si procede]','[Pregunta y necesidad]','[Autorización y costo separado]']], [2.3,2.5,2.2])
    box('Riesgos y alternativas.', '[Explicar qué evento cambiaría la ruta: notificación, plazo, nueva prueba, riesgo, contrapropuesta o imposibilidad de cumplimiento.]')
    page('5. Tiempos y condiciones')
    table(['Actividad','Tipo de tiempo','Inicio / rango','Fuente / contingencias'],[['[Actividad]','[Legal / interno / externo]','[Evento y plazo revisado]','[Base y variables]'],['[Etapa]','[Tipo]','[Rango sustentado]','[Dependencias]']], [1.5,1.5,2.0,2.0])
    text('Los rangos de planeación no constituyen promesa de duración. Los plazos legales deben indicar fundamento, evento inicial y calendario; las fechas de autoridad requieren comprobación.',True)
    doc.add_heading('6. Propuesta de honorarios',1)
    table(['Opción / etapa','Incluye y excluye','Monto e impuestos','Pago / hito'],[['[Opción A]','[Entregables y límites]','[Importe aprobado / moneda]','[Anticipo y saldo pactados]'],['[Opción B]','[Alcance distinto]','[Importe aprobado]','[Calendario]']], [1.2,2.3,1.7,1.8])
    box('Condiciones económicas.', '[Precisar vigencia, gastos de terceros, revisiones/sesiones incluidas, trabajo previo acreditable sólo si se acuerda y autorización necesaria para extras. Comprobar sumas y no duplicar conceptos.]')
    page('7. Próximos pasos y revisión')
    table(['Acción','Responsable','Plazo / condición','Comprobación'],[['[Acción prioritaria]','[Cliente / firma]','[Fecha real o condición]','[Documento o decisión]'],['[Siguiente acción]','[Responsable]','[Plazo]','[Evidencia]']], [2.1,1.3,1.8,1.8])
    doc.add_heading('Fuentes y límites',2)
    text('[Referencias verificadas y localizadores. Resumir los pendientes materiales y explicar cuándo debe actualizarse el análisis.]')
    box('Revisión profesional.', 'El despacho compromete preparación y diligencia dentro del alcance acordado; no garantiza decisiones de autoridades o terceros. La versión para cliente requiere revisión, datos correctos y aprobación de la persona abogada responsable.')
    text('Responsable: [Nombre y habilitación verificados]\nVersión y fecha real de aprobación: [Datos]\nFirma: [Mecanismo legítimo, no insertar por cuenta de la IA]')
    text('La plantilla no acredita asesoría prestada, contratación, pago, autorización o presentación de actuaciones.',True)
    text('Origen: MILLA Asesorías · lianfixx y colaboradores · Apache-2.0. Conservar LICENSE y NOTICE al redistribuir la plantilla. Otros profesionales deben adaptar membrete, folios y datos; no implica afiliación ni aval de MILLA ABOGADOS.', True)
    output.parent.mkdir(parents=True,exist_ok=True)
    if output.exists(): raise FileExistsError('La salida ya existe; elige otro nombre para no sobrescribirla.')
    doc.save(output)


def main():
    parser=argparse.ArgumentParser(description=__doc__); parser.add_argument('--output',type=Path,required=True); args=parser.parse_args()
    try: generate(args.output)
    except (OSError,RuntimeError) as exc: parser.exit(1,f'Error: {exc}\n')
    print(args.output)


if __name__=='__main__': main()
